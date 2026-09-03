import json
import os

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit,
    QPushButton, QTableWidget, QTableWidgetItem, QHeaderView,
    QAbstractItemView, QStyledItemDelegate, QMessageBox, QCompleter, QFileDialog
)
from PySide6.QtCore import Qt, QSize
from PySide6.QtGui import QGuiApplication


class EditorDelegate(QStyledItemDelegate):
    def sizeHint(self, option, index):
        base = super().sizeHint(option, index)
        return QSize(base.width(), max(base.height(), 34))

    def createEditor(self, parent, option, index):
        editor = super().createEditor(parent, option, index)
        if isinstance(editor, QLineEdit):
            editor.setMinimumHeight(34)
            editor.setStyleSheet("padding: 4px 8px;")
        return editor


def _caminho_json():
    import config
    base = config.obter_caminho_jsons()
    if not base:
        return ""
    return os.path.normpath(os.path.join(base, "Almox", "ControlePedidos", "FollowUp", "FollowUp.json"))


class FollowUpPage(QWidget):
    COLUNAS = [
        "Fornecedor", "N pedido", "DPP", "Kardex", "Código",
        "Qtde programada",
        "Qtde entregue", "Qtde pendente",
    ]

    def __init__(self):
        super().__init__()
        self.dados = []
        self._setup_ui()
        self._carregar_dados()

    def _setup_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(32, 32, 32, 32)
        layout.setSpacing(16)

        card = QWidget()
        card.setObjectName("pageCard")
        card_layout = QVBoxLayout(card)
        card_layout.setContentsMargins(28, 28, 28, 28)
        card_layout.setSpacing(16)

        titulo = QLabel("Follow-up")
        titulo.setObjectName("pageTitle")
        card_layout.addWidget(titulo)

        linha_top = QHBoxLayout()
        linha_top.setSpacing(8)

        self.campo_filtro = QLineEdit()
        self.campo_filtro.setPlaceholderText("Pesquisar...")
        self.campo_filtro.setFixedHeight(30)
        self.campo_filtro.setFixedWidth(200)
        self.campo_filtro.textChanged.connect(self._aplicar_filtro)
        linha_top.addWidget(self.campo_filtro)

        self.btn_exportar = QPushButton("Exportar Excel")
        self.btn_exportar.setFixedHeight(30)
        self.btn_exportar.clicked.connect(self._exportar_excel)
        linha_top.addWidget(self.btn_exportar)

        self.label_contador = QLabel("0 itens")
        self.label_contador.setObjectName("statusLabel")
        linha_top.addWidget(self.label_contador)

        linha_top.addStretch()
        card_layout.addLayout(linha_top)

        self.tabela = QTableWidget(0, len(self.COLUNAS))
        self.tabela.setHorizontalHeaderLabels(self.COLUNAS)
        header = self.tabela.horizontalHeader()
        header.setStyleSheet(
            "QHeaderView::section {"
            "  font-size: 9px;"
            "  font-weight: 600;"
            "  background-color: #f8fafc;"
            "  color: #64748b;"
            "  text-transform: uppercase;"
            "  letter-spacing: 0.5px;"
            "  padding: 8px 12px;"
            "  border: none;"
            "  border-bottom: 2px solid #e2e8f0;"
            "}"
        )
        header.setStretchLastSection(False)
        for c in range(self.tabela.columnCount()):
            header.setSectionResizeMode(c, QHeaderView.ResizeMode.Interactive)
            nome_col = self.COLUNAS[c]
            if nome_col == "Fornecedor": largura = 180
            elif nome_col == "N pedido": largura = 80
            elif nome_col == "DPP": largura = 80
            elif nome_col == "Kardex": largura = 150
            elif nome_col == "Código": largura = 180
            else: largura = 100
            header.resizeSection(c, largura)
        self.tabela.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectItems)
        self.tabela.setEditTriggers(QAbstractItemView.EditTrigger.DoubleClicked)
        self.tabela.setAlternatingRowColors(False)
        self.tabela.verticalHeader().setDefaultSectionSize(28)
        self.tabela.verticalHeader().setMinimumSectionSize(24)
        self.tabela.verticalHeader().setVisible(False)
        self.tabela.setHorizontalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)
        self.tabela.setItemDelegate(EditorDelegate(self.tabela))
        self.tabela.setVisible(True)
        card_layout.addWidget(self.tabela)

        layout.addWidget(card)

    def _popular_tabela(self):
        self.tabela.setVisible(True)
        self.tabela.blockSignals(True)
        self.tabela.setRowCount(0)
        filtro_texto = self.campo_filtro.text().strip().lower()
        
        import re
        def _sort_key(item):
            n_pedido = str(item.get("N do pedido", ""))
            return [int(c) if c.isdigit() else c.lower() for c in re.split(r'(\d+)', n_pedido)]
            
        dados_ordenados = sorted(self.dados, key=_sort_key)
        
        for item in dados_ordenados:
            if filtro_texto:
                texto = " ".join(str(v) for v in item.values()).lower()
                if filtro_texto not in texto:
                    continue
            row = self.tabela.rowCount()
            self.tabela.insertRow(row)
            for col, chave in enumerate(self.COLUNAS):
                chave_dado = "N do pedido" if chave == "N pedido" else chave
                valor = str(item.get(chave_dado, ""))
                cell = QTableWidgetItem(valor)
                cell.setFlags(cell.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self.tabela.setItem(row, col, cell)
        self.tabela.blockSignals(False)
        self.label_contador.setText(f"{self.tabela.rowCount()} itens")

    def atualizar_dados(self, dados):
        import copy
        self.dados = copy.deepcopy(dados)
        
        try:
            main_win = self.window()
            if main_win and hasattr(main_win, "pages"):
                ctrl_page = main_win.pages.get("controle_pedidos")
                if ctrl_page and hasattr(ctrl_page, "dados"):
                    dpp_map = {}
                    for item in ctrl_page.dados:
                        nome = item.get("nome", "").strip()
                        if nome:
                            pc = nome.split(" ")[0].strip()
                            dpp = item.get("dpp", "").strip()
                            if dpp:
                                dpp_map[pc] = dpp
                    
                    for item in self.dados:
                        n_pedido = item.get("N do pedido", "").strip()
                        if n_pedido and n_pedido in dpp_map:
                            item["DPP"] = dpp_map[n_pedido]
        except Exception:
            pass

        self._atualizar_completer()
        self._popular_tabela()
        self._salvar_follow_up()

    def _salvar_follow_up(self):
        import config
        base = config.obter_caminho_jsons()
        if not base:
            config.avisar_sem_pasta(self)
            return
        pasta_json = os.path.normpath(os.path.join(base, "Almox", "ControlePedidos", "FollowUp"))
        os.makedirs(pasta_json, exist_ok=True)
        caminho_arquivo = os.path.join(pasta_json, "FollowUp.json")
        with open(caminho_arquivo, "w", encoding="utf-8") as f:
            json.dump(self.dados, f, ensure_ascii=False, indent=2)

    def _carregar_dados(self):
        try:
            caminho = _caminho_json()
            if caminho and os.path.isfile(caminho):
                with open(caminho, "r", encoding="utf-8") as f:
                    self.dados = json.load(f)
            else:
                self.dados = []
        except (FileNotFoundError, json.JSONDecodeError, OSError):
            self.dados = []
        self.tabela.setVisible(True)
        self._atualizar_completer()
        self._popular_tabela()

    def _aplicar_filtro(self):
        self.tabela.setVisible(True)
        self._popular_tabela()

    def _atualizar_completer(self):
        fornecedores = sorted(list(set(
            str(item.get("Fornecedor", "")).strip() for item in self.dados if str(item.get("Fornecedor", "")).strip()
        )))
        completer = QCompleter(fornecedores, self)
        completer.setCaseSensitivity(Qt.CaseSensitivity.CaseInsensitive)
        completer.setFilterMode(Qt.MatchFlag.MatchContains)
        self.campo_filtro.setCompleter(completer)

    def _exportar_excel(self):
        try:
            import pandas as pd
        except ImportError:
            QMessageBox.warning(self, "Aviso", "Biblioteca pandas não encontrada. Instale-a para exportar.")
            return

        colunas_desejadas = ["DPP", "Kardex", "Código", "Qtde programada", "Qtde entregue", "Qtde pendente"]
        colunas_exportacao = ["DPP", "KARDEX", "CÓDIGO", "PROGRAMADO", "ENTREGUE", "PENDENTE"]
        colunas_indices = []
        for col in colunas_desejadas:
            if col in self.COLUNAS:
                colunas_indices.append(self.COLUNAS.index(col))
            else:
                colunas_indices.append(-1)

        dados_exportar = []
        for row in range(self.tabela.rowCount()):
            linha_dados = []
            for col_idx in colunas_indices:
                if col_idx != -1:
                    item = self.tabela.item(row, col_idx)
                    linha_dados.append(item.text() if item else "")
                else:
                    linha_dados.append("")
            dados_exportar.append(linha_dados)

        df = pd.DataFrame(dados_exportar, columns=colunas_exportacao)

        caminho, _ = QFileDialog.getSaveFileName(self, "Exportar para Excel", "", "Excel Files (*.xlsx)")
        if caminho:
            if not caminho.endswith(".xlsx"):
                caminho += ".xlsx"
            try:
                from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
                
                with pd.ExcelWriter(caminho, engine='openpyxl') as writer:
                    df.to_excel(writer, index=False)
                    worksheet = writer.sheets['Sheet1']
                    
                    header_fill = PatternFill(start_color="0070C0", end_color="0070C0", fill_type="solid") # Azul
                    header_font = Font(color="FFFFFF", bold=True)
                    center_alignment = Alignment(horizontal="center", vertical="center")
                    thin_border = Border(
                        left=Side(style='thin'), 
                        right=Side(style='thin'), 
                        top=Side(style='thin'), 
                        bottom=Side(style='thin')
                    )

                    for col_num, value in enumerate(df.columns.values):
                        cell = worksheet.cell(row=1, column=col_num + 1)
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = center_alignment
                        cell.border = thin_border
                        # Ajustar largura da coluna
                        column_letter = cell.column_letter
                        worksheet.column_dimensions[column_letter].width = 15

                    for row_num in range(len(df)):
                        for col_num in range(len(df.columns)):
                            cell = worksheet.cell(row=row_num + 2, column=col_num + 1)
                            cell.alignment = center_alignment
                            cell.border = thin_border
                            
                QMessageBox.information(self, "Sucesso", "Dados exportados com sucesso!")
            except Exception as e:
                QMessageBox.critical(self, "Erro", f"Ocorreu um erro ao exportar: {str(e)}")
